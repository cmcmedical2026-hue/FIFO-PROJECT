"""Excel workbook for a lot ledger result (scripts/lot_ledger.py --xlsx).

Kept apart from flow.lots so the ledger itself stays standard-library only;
this module is the one place that needs openpyxl.
"""

from __future__ import annotations

import datetime as dt
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OP_NAMES = {
    "STROUT": "إذن صرف", "STRIN2": "إذن إضافة", "STRMRT": "ارتجاع مبيعات",
    "STRIN9": "تحويل داخل", "STROU9": "تحويل خارج", "BURRET": "مرتجع مورد",
    "STRMAK": "صرف عينات", "STRMR": "ارتجاع عينات", "STRROT": "تسوية جردية مدينة",
    "STRRTS": "تسوية جردية دائنة", "RELABEL": "تعديل لوط على الجرد", "COST-ADJ": "تسوية تكلفة",
}
HOW = {"remark": "من الملاحظة", "override": "قرار/تقسيم يدوي", "remark-split": "من الملاحظة (مقسم)",
       "fefo": "مقترح: الأقرب انتهاءً", "default": "مقترح: اللوط الوحيد/الأول", "count": "من الجرد",
       "settle": "تسوية"}

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
GREY = PatternFill("solid", fgColor="EDEDED")
RED = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFE699")
GREEN = PatternFill("solid", fgColor="C6EFCE")


def _expiry_date(exp: str) -> dt.date | None:
    parts = [int(p) for p in exp.replace("/", "-").split("-") if p.isdigit()]
    try:
        if len(parts) == 3:
            return dt.date(parts[2], parts[1], parts[0])
        if len(parts) == 2:
            nxt = dt.date(parts[1] + (parts[0] == 12), parts[0] % 12 + 1, 1)
            return nxt - dt.timedelta(days=1)
        if len(parts) == 1 and parts[0] > 2000:
            return dt.date(parts[0], 12, 31)
    except ValueError:
        return None
    return None


def _sheet(wb, title, header, rows, widths=None, fills=None, money=()):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for n, r in enumerate(rows):
        ws.append(r)
        if fills and fills[n]:
            for c in ws[ws.max_row]:
                c.fill = fills[n]
    for col in money:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row:
                c.number_format = "#,##0.00"
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def _key(item: str):
    return (not item.isdigit(), item.zfill(8))


def _known_sum(values):
    values = list(values)
    return None if any(v is None for v in values) else sum(values)


def write_xlsx(res: dict, path: str, *, snapshot_at: dict | None = None,
               questions: list[list[str]] | None = None, today: dt.date | None = None,
               opening_path: str | None = None) -> None:
    today = today or dt.date.today()
    names = res["names"]
    wb = Workbook()
    wb.remove(wb.active)

    # -- current balance by lot and batch -----------------------------------
    rows, fills = [], []
    for b in sorted(res["balance"], key=lambda r: (_key(r["item"]), r["lot"])):
        e = _expiry_date(b["exp"])
        note = ""
        fill = None
        if e and e < today:
            note, fill = "منتهي", RED
        elif e and (e - today).days <= 90:
            note, fill = "ينتهي خلال 3 شهور", AMBER
        if b["qty"] < 0:
            note, fill = "سالب — راجع", RED
        rows.append([b["item"], names.get(b["item"], ""), b["lot"], b["exp"], b["qty"],
                     b["cost"], b["value"], b["batch"], note])
        fills.append(fill)
    rows.append(["", "الإجمالي", "", "", sum(r[4] for r in rows), "", _known_sum(r[6] for r in rows), "", ""])
    fills.append(GREY)
    _sheet(wb, f"رصيد باللوط {res['to']}",
           ["كود الصنف", "اسم الصنف", "اللوط", "الصلاحية", "الكمية", "تكلفة الوحدة", "القيمة",
            "الدفعة", "تنبيه"], rows, [11, 38, 16, 12, 9, 12, 14, 34, 16], fills, money=(6, 7))

    # -- item summary ---------------------------------------------------------
    now = defaultdict(lambda: [0.0, 0.0, 0, False])
    for b in res["balance"]:
        now[b["item"]][0] += b["qty"]
        if b["value"] is None:
            now[b["item"]][3] = True
        elif not now[b["item"]][3]:
            now[b["item"]][1] += b["value"]
    lots_now = defaultdict(set)
    for b in res["balance"]:
        lots_now[b["item"]].add(b["lot"])
    snap = defaultdict(float)
    counted: dict[str, float] = {}
    if snapshot_at:
        for b in snapshot_at["balance"]:
            snap[b["item"]] += b["qty"]
        last = snapshot_at["checkpoints"][-1] if snapshot_at["checkpoints"] else None
        if last:
            data = json.loads(Path(last["count_path"]).read_text(encoding="utf-8"))
            for row in data["lines"]:
                counted[str(row[0])] = counted.get(str(row[0]), 0) + row[2]
    items = sorted(set(now) | set(snap) | set(counted), key=_key)
    rows, fills = [], []
    for i in items:
        c = counted.get(i, 0)
        s = snap.get(i, 0)
        diff = c - s
        rows.append([i, names.get(i, ""), s, c, diff, "✓" if abs(diff) < 1e-9 else "✗",
                     now[i][0], len([l for l in lots_now[i]]),
                     None if now[i][3] else round(now[i][1], 2)])
        fills.append(None if abs(diff) < 1e-9 else RED)
    rows.append(["", "الإجمالي", sum(r[2] for r in rows), sum(r[3] for r in rows),
                 sum(r[4] for r in rows), "", sum(r[6] for r in rows), "", _known_sum(r[8] for r in rows)])
    fills.append(GREY)
    cp_date = snapshot_at["to"] if snapshot_at else ""
    _sheet(wb, "ملخص بالصنف",
           ["كود الصنف", "اسم الصنف", f"رصيد فلو {cp_date}", f"الجرد الفعلي {cp_date}", "الفرق (جرد - فلو)",
            "مطابق؟", f"الرصيد {res['to']}", "عدد اللوطات", f"القيمة {res['to']}"],
           rows, [11, 38, 14, 14, 12, 8, 14, 10, 16], fills, money=(9,))

    # -- balance at the last count -------------------------------------------
    if snapshot_at:
        rows = [[b["item"], names.get(b["item"], ""), b["lot"], b["exp"], b["qty"], b["cost"],
                 b["value"], b["batch"]]
                for b in sorted(snapshot_at["balance"], key=lambda r: (_key(r["item"]), r["lot"]))]
        rows.append(["", "الإجمالي", "", "", sum(r[4] for r in rows), "", _known_sum(r[6] for r in rows), ""])
        _sheet(wb, f"رصيد باللوط {snapshot_at['to']}",
               ["كود الصنف", "اسم الصنف", "اللوط", "الصلاحية", "الكمية", "تكلفة الوحدة", "القيمة",
                "الدفعة"], rows, [11, 38, 16, 12, 9, 12, 14, 34], money=(6, 7))

    # -- count reconciliation -------------------------------------------------
    rows, fills = [], []
    for cp in res["checkpoints"]:
        for d in cp["diffs_before"]:
            rows.append([cp["as_of"], "فرق قبل التعديل", d["item"], names.get(d["item"], ""), d["lot"],
                         d["ledger"], d["count"], d["diff"],
                         "فرق لوط (الصنف مطابق)" if d["item_ledger"] == d["item_count"] else "فرق كمية في الصنف"])
            fills.append(None if d["item_ledger"] == d["item_count"] else RED)
        for m in cp["relabels"]:
            rows.append([cp["as_of"], "تعديل لوط", m["item"], names.get(m["item"], ""),
                         f"{m['from']} ← {m['to']}", "", "", m["qty"], "اتنقلت بتكلفة دفعتها — مفيش أثر على القيمة"])
            fills.append(GREEN)
        for d in cp["diffs_after"]:
            rows.append([cp["as_of"], "فرق مفتوح", d["item"], names.get(d["item"], ""), d["lot"],
                         d["ledger"], d["count"], d["diff"], "مستني قرار"])
            fills.append(AMBER)
    _sheet(wb, "مطابقة الجرد",
           ["تاريخ الجرد", "البند", "كود الصنف", "اسم الصنف", "اللوط", "دفتر اللوطات", "الجرد", "الفرق", "ملاحظة"],
           rows, [12, 16, 11, 34, 30, 12, 10, 10, 40], fills)

    # -- movement lines -------------------------------------------------------
    rows, fills = [], []
    for r in res["log"]:
        q = r["qty"]
        rows.append([r["date"], OP_NAMES.get(r["op"], r["op"]), r["doc"], r["item"], names.get(r["item"], ""),
                     r["lot"], r["exp"], q if r["dir"] == "A" else "", q if r["dir"] == "S" else "",
                     r["value"], HOW.get(r["how"], r["how"]), r["party"], r["remark"][:250]])
        fills.append(AMBER if r["how"] in ("fefo", "default") else
                     GREEN if r["op"] in ("RELABEL", "COST-ADJ") else None)
    _sheet(wb, "الحركات باللوط",
           ["التاريخ", "الحركة", "رقم المستند", "كود الصنف", "اسم الصنف", "اللوط", "الصلاحية", "وارد", "منصرف",
            "القيمة بالتكلفة", "مصدر اللوط", "الجهة", "الملاحظة في فلو"],
           rows, [11, 16, 9, 10, 30, 14, 11, 8, 8, 13, 20, 26, 60], fills, money=(10,))

    # -- opening --------------------------------------------------------------
    if opening_path:
        op = json.loads(Path(opening_path).read_text(encoding="utf-8"))
        rows = [[b[0], names.get(str(b[0]), ""), b[1], b[2], b[3], round(b[2] * b[3], 2), b[4]]
                for b in op["batches"]]
        rows.append(["", "الإجمالي", "", sum(r[3] for r in rows), "", sum(r[5] for r in rows), ""])
        _sheet(wb, f"رصيد أول المدة {op['date']}",
               ["كود الصنف", "اسم الصنف", "اللوط", "الكمية", "تكلفة الوحدة", "القيمة", "ملاحظة"],
               rows, [11, 38, 16, 9, 12, 14, 50], money=(5, 6))

    # -- questions and method -------------------------------------------------
    if questions:
        _sheet(wb, "أسئلة مفتوحة", ["#", "الصنف", "السؤال", "أثر على التكلفة؟", "الافتراض الحالي"],
               [[n + 1, *q] for n, q in enumerate(questions)], [4, 12, 80, 22, 50])
    method = [
        ["المخزن", "مخزون المقطم 1 (company 1 / store 1)"],
        ["نقطة البداية", "رصيد 02-05-2026 باللوط والتكلفة (الشيت اللي على الدرايف) — مطابق لفلو 36 صنف / 5869 قطعة"],
        ["الكمية", "من فلو: trndtl, STS='A', وارد EFFECTITBAL='A' / منصرف 'S'، والكمية STRQTY (بتشمل البونص)"],
        ["اللوط", "بالترتيب: قرار يدوي مسجل ← المكتوب في الـremark ← الأقرب انتهاءً (مقترح، ملون أصفر)"],
        ["الجرد", "كل جرد نقطة مطابقة: لو إجمالي الصنف مطابق بيتعدل توزيع اللوطات على الجرد (بتكلفة الدفعة نفسها). لو الإجمالي مختلف بيفضل فرق مفتوح لحد القرار"],
        ["جرد 01-07", "مطابق لفلو آخر يوم 30-06 (قبل تحويلات 01-07)؛ فسّر توزيع إذن الإضافة 125 على اللوطات"],
        ["التكلفة", "كل دفعة بسعر شرائها. الصرف من اللوط بياخد أقدم دفعة فيه. المرتجع/التحويل الراجع بتكلفة الدفعة اللي خرج منها. مفيش متوسط ومفيش تكلفة من فلو"],
        ["تسوية التكلفة", "لو لوط اتصرف منه أكتر من رصيده وبعدين الجرد بيّن إن الزيادة كانت من لوط تاني بسعر مختلف، الفرق بيتسجل سطر 'تسوية تكلفة'"],
    ]
    _sheet(wb, "طريقة الحساب", ["البند", "التفاصيل"], method, [18, 120])
    wb.save(path)
